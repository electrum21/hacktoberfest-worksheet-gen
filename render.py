import io
from html import escape

LETTERS = "ABCD"


def _br(s: str) -> str:
    return escape(s).replace("\n", "<br>")


def _meta(ws, topics, difficulty, tutor, student):
    total = sum(q.marks for q in ws.questions)
    bits = [f"Difficulty: {difficulty}", f"Total: {total} marks",
            f"Suggested time (rough guide): {round(total * 1.2)} min"]
    if student:
        bits.append(f"For: {student}")
    if tutor:
        bits.append(tutor)
    return total, " · ".join(bits)


def to_html(ws, subject, topics, difficulty, tutor="", student="", include_key=True) -> str:
    total, meta = _meta(ws, topics, difficulty, tutor, student)
    items, key = [], []
    for q in ws.questions:
        opts = ""
        if q.qtype == "mcq":
            opts = '<ol class="opts" type="A">' + "".join(f"<li>{escape(o)}</li>" for o in q.options) + "</ol>"
        space = "" if q.qtype == "mcq" else f'<div style="height:{min(40 * q.marks, 200)}px"></div>'
        items.append(f'<li><div>{_br(q.question)} <span class="m">[{q.marks}m]</span></div>{opts}{space}</li>')
        ms = ""
        if q.mark_scheme:
            ms = '<ul class="ms">' + "".join(f"<li>{escape(m)}</li>" for m in q.mark_scheme) + "</ul>"
        key.append(f'<li>{_br(q.worked_solution)}{ms}<div class="a">Answer: {escape(q.final_answer)}</div></li>')
    key_html = (f'<div class="key"><h1>Answer Key and Mark Scheme</h1><ol>{"".join(key)}</ol></div>'
                if include_key else "")
    return f"""<!doctype html><html><head><meta charset="utf-8">
<title>{escape(subject)}</title>
<style>
 body{{font-family:Georgia,serif;max-width:760px;margin:2rem auto;padding:0 1rem;line-height:1.5}}
 h1{{font-size:1.4rem;margin-bottom:.2rem}} .meta{{color:#555;font-size:.9rem}}
 .hdr{{display:flex;gap:2rem;margin:1rem 0;border-bottom:1px solid #000;padding-bottom:.5rem}}
 li{{margin-bottom:1rem}} .m{{color:#666;font-size:.85rem}} .opts li{{margin:0}}
 .ms{{color:#333;font-size:.9rem}} .a{{margin-top:.3rem;font-weight:bold}}
 .key{{page-break-before:always}} @media print{{body{{margin:0}}}}
</style></head><body>
<h1>{escape(subject)}: {escape(topics)}</h1>
<div class="meta">{escape(meta)}</div>
<div class="hdr"><span>Name: ____________________</span><span>Date: ___________</span><span>Score: ____ / {total}</span></div>
<ol>{''.join(items)}</ol>
{key_html}
</body></html>"""


def to_docx(ws, subject, topics, difficulty, tutor="", student="") -> bytes:
    from docx import Document
    from docx.shared import Cm

    total, meta = _meta(ws, topics, difficulty, tutor, student)
    d = Document()
    d.add_heading(f"{subject}: {topics}", level=1)
    d.add_paragraph(meta)
    d.add_paragraph(f"Name: ____________________    Date: ___________    Score: ____ / {total}")
    for i, q in enumerate(ws.questions, 1):
        p = d.add_paragraph()
        p.add_run(f"{i}. ").bold = True
        p.add_run(f"{q.question}  [{q.marks}m]")
        if q.qtype == "mcq":
            for letter, o in zip(LETTERS, q.options):
                d.add_paragraph(f"{letter}. {o}").paragraph_format.left_indent = Cm(1)
        else:
            for _ in range(min(q.marks, 4) + 1):
                d.add_paragraph("")
    d.add_page_break()
    d.add_heading("Answer Key and Mark Scheme", level=1)
    for i, q in enumerate(ws.questions, 1):
        p = d.add_paragraph()
        p.add_run(f"{i}. Answer: {q.final_answer}").bold = True
        if q.worked_solution:
            d.add_paragraph(q.worked_solution)
        for m in q.mark_scheme:
            d.add_paragraph(m, style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
